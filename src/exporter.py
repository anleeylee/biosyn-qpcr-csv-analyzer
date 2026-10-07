"""Export qPCR results to XLSX / HTML."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import List, Dict, Any

from .analyzer import (
    qPCRRecord, CleanRecord, ReplicateStats,
    DCtResult, DDCtResult,
)

logger = logging.getLogger(__name__)


def export_qpcr_results(
    raw_records: List[qPCRRecord],
    cleaned: List[CleanRecord],
    rep_stats: List[ReplicateStats],
    dct_results: List[DCtResult],
    ddct_results: List[DDCtResult],
    exclusions: List[Dict],
    output_path: Path,
) -> None:
    """Export all results to XLSX (default) or HTML."""
    suffix = output_path.suffix.lower()
    if suffix == ".html":
        _export_html(rep_stats, ddct_results, output_path)
    else:
        _export_xlsx(
            raw_records, cleaned, rep_stats,
            dct_results, ddct_results, exclusions, output_path,
        )
    logger.info("Exported qPCR results to %s", output_path)


def _export_xlsx(
    raw: List[qPCRRecord],
    cleaned: List[CleanRecord],
    rep_stats: List[ReplicateStats],
    dct: List[DCtResult],
    ddct: List[DDCtResult],
    exclusions: List[Dict],
    path: Path,
) -> None:
    """Export to multi-sheet XLSX."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    hdr = Font(bold=True)

    # Raw Data
    ws = wb.active
    ws.title = "Raw Data"
    ws.append(["Sample", "Group", "Gene", "Well", "Ct", "Replicate"])
    for c in ws[1]:
        c.font = hdr
    for r in raw:
        ws.append([r.sample, r.group, r.gene, r.well, r.ct, r.replicate])

    # Clean Data
    ws2 = wb.create_sheet("Clean Data")
    ws2.append(["Sample", "Group", "Gene", "Well", "Ct", "Replicate", "Excluded", "Reason"])
    for c in ws2[1]:
        c.font = hdr
    for cr in cleaned:
        r = cr.record
        ws2.append([r.sample, r.group, r.gene, r.well, r.ct, r.replicate,
                    cr.excluded, cr.exclusion_reason])

    # Replicate QC
    ws3 = wb.create_sheet("Replicate QC")
    ws3.append(["Sample", "Group", "Gene", "n", "Mean Ct", "SD", "CV%",
                "Min", "Max", "Range", "Status"])
    for c in ws3[1]:
        c.font = hdr
    for rs in rep_stats:
        ws3.append([rs.sample, rs.group, rs.gene, rs.n, rs.mean_ct, rs.sd,
                    rs.cv, rs.min_ct, rs.max_ct, rs.range_ct, rs.status])

    # ΔCt
    ws4 = wb.create_sheet("Delta Ct")
    ws4.append(["Sample", "Group", "Target", "Reference",
                "Target Ct", "Ref Ct", "ΔCt"])
    for c in ws4[1]:
        c.font = hdr
    for d in dct:
        ws4.append([d.sample, d.group, d.target_gene, d.reference_gene,
                    d.mean_target_ct, d.mean_ref_ct, d.delta_ct])

    # ΔΔCt & Fold Change
    ws5 = wb.create_sheet("Fold Change")
    ws5.append(["Sample", "Group", "Gene", "ΔCt", "ΔΔCt",
                "Fold Change (2^-ΔΔCt)", "Is Calibrator"])
    for c in ws5[1]:
        c.font = hdr
    for d in ddct:
        ws5.append([d.sample, d.group, d.gene, d.delta_ct,
                    d.delta_delta_ct, d.fold_change, d.is_calibrator])

    # Errors / Exclusions
    ws6 = wb.create_sheet("Errors")
    ws6.append(["Sample", "Gene", "Well", "Reason"])
    for c in ws6[1]:
        c.font = hdr
    for e in exclusions:
        ws6.append([e["sample"], e["gene"], e["well"], e["reason"]])

    # Methods
    ws7 = wb.create_sheet("Methods")
    methods = [
        ("Method", "Description"),
        ("ΔCt", "Ct(target) - Ct(reference)"),
        ("ΔΔCt", "ΔCt(sample) - ΔCt(calibrator)"),
        ("Fold Change", "2^(-ΔΔCt) — assumes ~100% amplification efficiency"),
        ("Replicate QC", "CV% threshold configurable; Z-score outlier detection"),
        ("", ""),
        ("Disclaimer", "Computational results based on supplied data. "
                       "Does not replace experimental QC. "
                       "Does not prove biological significance."),
        ("Source", "BioSyn qPCR CSV Analyzer v1.0.0 — https://biosyn-inc.com/"),
    ]
    for row in methods:
        ws7.append(row)

    wb.save(path)


def _export_html(
    rep_stats: List[ReplicateStats],
    ddct: List[DDCtResult],
    path: Path,
) -> None:
    """Export standalone HTML report."""
    rep_rows = ""
    for rs in rep_stats:
        color = {"PASS": "#28a745", "REVIEW": "#ffc107",
                 "OUTLIER": "#dc3545"}.get(rs.status, "#6c757d")
        rep_rows += f"""<tr><td>{rs.sample}</td><td>{rs.gene}</td>
        <td>{rs.n}</td><td>{rs.mean_ct:.3f}</td><td>{rs.sd:.3f}</td>
        <td>{rs.cv:.2f}%</td>
        <td style="color:{color};font-weight:bold">{rs.status}</td></tr>"""

    fc_rows = ""
    for d in ddct:
        fc_rows += f"""<tr><td>{d.sample}</td><td>{d.group}</td>
        <td>{d.gene}</td><td>{d.delta_ct:.3f}</td>
        <td>{d.delta_delta_ct:.3f}</td><td>{d.fold_change:.4f}</td></tr>"""

    html = f"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<title>BioSyn qPCR Analysis Report</title>
<style>
body{{font-family:system-ui;margin:2rem;background:#f8f9fa}}
table{{border-collapse:collapse;width:100%;background:#fff;margin-bottom:2rem}}
th,td{{padding:.5rem 1rem;border:1px solid #dee2e6;text-align:left}}
th{{background:#e9ecef}}
</style></head><body>
<h1>BioSyn qPCR Analysis Report</h1>
<p>By <a href="https://biosyn-inc.com/">BioSyn Inc.</a> — local-first qPCR data analysis.</p>
<h2>Replicate QC</h2>
<table><tr><th>Sample</th><th>Gene</th><th>n</th><th>Mean Ct</th>
<th>SD</th><th>CV%</th><th>Status</th></tr>{rep_rows}</table>
<h2>Fold Change (2^-ΔΔCt)</h2>
<table><tr><th>Sample</th><th>Group</th><th>Gene</th>
<th>ΔCt</th><th>ΔΔCt</th><th>Fold Change</th></tr>{fc_rows}</table>
<p style="font-size:.85rem;color:#6c757d;margin-top:2rem">
Computational results based on supplied data. Does not replace experimental QC.
See <a href="https://biosyn-inc.com/tools/qpcr-analysis">BioSyn qPCR Analysis Tools</a> for more tools.</p>
</body></html>"""
    path.write_text(html, encoding="utf-8")
